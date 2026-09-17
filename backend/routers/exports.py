import csv
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, HTTPException, Response, status
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.database import get_db
from backend.models import (
    FieldOption,
    Form,
    FormField,
    FormVersion,
    Survey,
    SurveyAnswer,
    User,
)


router = APIRouter(
    prefix="/exports",
    tags=["exports"],
)


# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def get_answer_export_value(
    answer: SurveyAnswer | None,
    option: FieldOption | None,
):
    """
    Convierte una SurveyAnswer en un valor fácil de exportar.
    Se utiliza principalmente para CSV.
    """

    if answer is None:
        return ""

    # SELECT
    if answer.field_option_id is not None:
        if option is not None:
            return option.label

        return str(answer.field_option_id)

    # TEXT / TEXTAREA
    if answer.value_text is not None:
        return answer.value_text

    # NUMBER
    if answer.value_number is not None:
        return str(answer.value_number)

    # DATE
    if answer.value_date is not None:
        return answer.value_date.isoformat()

    # BOOLEAN
    if answer.value_boolean is not None:
        return (
            "true"
            if answer.value_boolean
            else "false"
        )

    return ""


def get_export_data(
    form_version_id: int,
    db: Session,
):
    """
    Obtiene todos los datos necesarios para exportar
    una versión de formulario.

    Esta función es compartida por CSV y Excel.
    """

    # 1. Buscar versión
    version = db.scalar(
        select(FormVersion).where(
            FormVersion.id == form_version_id
        )
    )

    if version is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Versión de formulario no encontrada",
        )

    # 2. Buscar formulario
    form = db.scalar(
        select(Form).where(
            Form.id == version.form_id
        )
    )

    if form is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Formulario no encontrado",
        )

    # 3. Campos de la versión
    fields = db.scalars(
        select(FormField)
        .where(
            FormField.form_version_id
            == form_version_id
        )
        .order_by(
            FormField.field_order
        )
    ).all()

    # 4. Encuestas de esa versión
    survey_rows = db.execute(
        select(
            Survey,
            User,
        )
        .join(
            User,
            Survey.user_id == User.id,
        )
        .where(
            Survey.form_version_id
            == form_version_id,
            Survey.status == "submitted",
        )
        .order_by(
            Survey.received_at
        )
    ).all()

    survey_ids = [
        survey.id
        for survey, user in survey_rows
    ]

    # 5. Respuestas agrupadas
    answers_by_survey_and_field = {}

    if survey_ids:
        answer_rows = db.execute(
            select(
                SurveyAnswer,
                FieldOption,
            )
            .outerjoin(
                FieldOption,
                SurveyAnswer.field_option_id
                == FieldOption.id,
            )
            .where(
                SurveyAnswer.survey_id.in_(
                    survey_ids
                )
            )
        ).all()

        for answer, option in answer_rows:
            answers_by_survey_and_field[
                (
                    answer.survey_id,
                    answer.form_field_id,
                )
            ] = (
                answer,
                option,
            )

    return (
        version,
        form,
        fields,
        survey_rows,
        answers_by_survey_and_field,
    )


# ============================================================
# EXPORTACIÓN CSV
# ============================================================

@router.get(
    "/form-versions/{form_version_id}/csv",
    responses={
        200: {
            "description": (
                "Archivo CSV generado correctamente"
            )
        },
        404: {
            "description": (
                "Versión o formulario no encontrado"
            )
        },
    },
)
def export_form_version_csv(
    form_version_id: int,
    db: Session = Depends(get_db),
):
    (
        version,
        form,
        fields,
        survey_rows,
        answers_by_survey_and_field,
    ) = get_export_data(
        form_version_id,
        db,
    )

    # Crear CSV en memoria
    output = StringIO(
        newline=""
    )

    writer = csv.writer(
        output,
        delimiter=";",
        lineterminator="\n",
    )

    # Columnas generales
    headers = [
        "survey_id",
        "uuid",
        "project_id",
        "form_id",
        "form_version_id",
        "version_number",
        "user_id",
        "employee_number",
        "status",
        "captured_at",
        "received_at",
        "latitude",
        "longitude",
    ]

    # Cada pregunta se convierte en columna
    field_headers = [
        f"field_{field.id}_{field.name}"
        for field in fields
    ]

    writer.writerow(
        headers + field_headers
    )

    # Una fila por encuesta
    for survey, user in survey_rows:
        row = [
            survey.id,
            str(survey.uuid),
            form.project_id,
            form.id,
            version.id,
            version.version_number,
            user.id,
            user.employee_number,
            survey.status,
            survey.captured_at.isoformat(),
            survey.received_at.isoformat(),

            # En CSV usamos coma decimal
            # para facilitar apertura directa
            # con Excel en configuración colombiana.
            (
                str(
                    survey.latitude
                ).replace(".", ",")
                if survey.latitude
                is not None
                else ""
            ),

            (
                str(
                    survey.longitude
                ).replace(".", ",")
                if survey.longitude
                is not None
                else ""
            ),
        ]

        field_values = []

        for field in fields:
            answer_data = (
                answers_by_survey_and_field.get(
                    (
                        survey.id,
                        field.id,
                    )
                )
            )

            if answer_data is None:
                field_values.append("")
                continue

            answer, option = answer_data

            field_values.append(
                get_answer_export_value(
                    answer,
                    option,
                )
            )

        writer.writerow(
            row + field_values
        )

    # UTF-8 + BOM
    # Esto permite que Excel reconozca
    # correctamente tildes, ñ, etc.
    csv_content = (
        "\ufeff"
        + output.getvalue()
    ).encode("utf-8")

    filename = (
        f"form_{form.id}"
        f"_version_{version.version_number}"
        ".csv"
    )

    return Response(
        content=csv_content,
        media_type=(
            "text/csv; charset=utf-8"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; '
                f'filename="{filename}"'
            )
        },
    )


# ============================================================
# EXPORTACIÓN EXCEL XLSX
# ============================================================

@router.get(
    "/form-versions/{form_version_id}/xlsx",
    responses={
        200: {
            "description": (
                "Archivo Excel generado correctamente"
            )
        },
        404: {
            "description": (
                "Versión o formulario no encontrado"
            )
        },
    },
)
def export_form_version_xlsx(
    form_version_id: int,
    db: Session = Depends(get_db),
):
    (
        version,
        form,
        fields,
        survey_rows,
        answers_by_survey_and_field,
    ) = get_export_data(
        form_version_id,
        db,
    )

    # Crear libro Excel
    workbook = Workbook()

    # Workbook() crea automáticamente
    # una hoja real tipo Worksheet.
    worksheet = workbook.worksheets[0]

    worksheet.title = "Encuestas"

    # --------------------------------------------------------
    # Encabezados
    # --------------------------------------------------------

    headers = [
        "survey_id",
        "uuid",
        "project_id",
        "form_id",
        "form_version_id",
        "version_number",
        "user_id",
        "employee_number",
        "status",
        "captured_at",
        "received_at",
        "latitude",
        "longitude",
    ]

    field_headers = [
        f"field_{field.id}_{field.name}"
        for field in fields
    ]

    worksheet.append(
        headers + field_headers
    )

    # Negrita para encabezados
    for cell in worksheet[1]:
        cell.font = Font(
            bold=True
        )

    # --------------------------------------------------------
    # Filas
    # --------------------------------------------------------

    for survey, user in survey_rows:
        row = [
            survey.id,
            str(survey.uuid),
            form.project_id,
            form.id,
            version.id,
            version.version_number,
            user.id,
            user.employee_number,
            survey.status,

            # En XLSX se conservan como
            # datetime reales.
            survey.captured_at,
            survey.received_at,

            # En XLSX mantenemos
            # coordenadas como números.
            survey.latitude,
            survey.longitude,
        ]

        field_values = []

        for field in fields:
            answer_data = (
                answers_by_survey_and_field.get(
                    (
                        survey.id,
                        field.id,
                    )
                )
            )

            if answer_data is None:
                field_values.append(
                    None
                )
                continue

            answer, option = answer_data

            value = None

            # SELECT
            if (
                answer.field_option_id
                is not None
            ):
                if option is not None:
                    value = option.label
                else:
                    value = (
                        answer.field_option_id
                    )

            # TEXT / TEXTAREA
            elif (
                answer.value_text
                is not None
            ):
                value = (
                    answer.value_text
                )

            # NUMBER
            elif (
                answer.value_number
                is not None
            ):
                value = float(
                    answer.value_number
                )

            # DATE
            elif (
                answer.value_date
                is not None
            ):
                value = (
                    answer.value_date
                )

            # BOOLEAN
            elif (
                answer.value_boolean
                is not None
            ):
                value = (
                    answer.value_boolean
                )

            field_values.append(
                value
            )

        worksheet.append(
            row + field_values
        )

    # --------------------------------------------------------
    # CONFIGURACIÓN VISUAL DEL EXCEL
    # --------------------------------------------------------

    # Congelar encabezado
    worksheet.freeze_panes = "A2"

    # Activar filtros
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # Ajustar ancho de columnas
    for column_number in range(
        1,
        worksheet.max_column + 1,
    ):
        max_length = 0

        column_letter = (
            get_column_letter(
                column_number
            )
        )

        for row_number in range(
            1,
            worksheet.max_row + 1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            if cell.value is None:
                continue

            value_length = len(
                str(
                    cell.value
                )
            )

            if (
                value_length
                > max_length
            ):
                max_length = (
                    value_length
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40,
        )

    # --------------------------------------------------------
    # FORMATO DE COORDENADAS
    # --------------------------------------------------------

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        # latitude = columna 12
        worksheet.cell(
            row=row_number,
            column=12,
        ).number_format = "0.000000"

        # longitude = columna 13
        worksheet.cell(
            row=row_number,
            column=13,
        ).number_format = "0.000000"

    # --------------------------------------------------------
    # FORMATO DE FECHAS
    # --------------------------------------------------------

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        # captured_at
        worksheet.cell(
            row=row_number,
            column=10,
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        # received_at
        worksheet.cell(
            row=row_number,
            column=11,
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

    # --------------------------------------------------------
    # GUARDAR ARCHIVO EN MEMORIA
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(
        output
    )

    excel_content = (
        output.getvalue()
    )

    filename = (
        f"form_{form.id}"
        f"_version_{version.version_number}"
        ".xlsx"
    )

    return Response(
        content=excel_content,
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; '
                f'filename="{filename}"'
            )
        },
    )